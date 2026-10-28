import zipfile
from pathlib import Path

import openpyxl
import pytest

from entity_resolver import tabular


def write_xlsx(path: Path, rows: list[list], sheet: str = "Data") -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for row in rows:
        ws.append(row)
    wb.save(path)
    return path


def test_detect_kind_accepts_matching_types(tmp_path):
    x = write_xlsx(tmp_path / "a.xlsx", [["a"]])
    c = tmp_path / "a.csv"
    c.write_text("a,b\n1,2\n")
    assert tabular.detect_kind(x, "a.xlsx") == "xlsx"
    assert tabular.detect_kind(c, "a.csv") == "csv"


@pytest.mark.parametrize("name", ["a.xlsm", "a.exe", "a.txt", "noext", "a.xlsx.exe"])
def test_detect_kind_rejects_other_extensions(tmp_path, name):
    p = tmp_path / "f"
    p.write_bytes(b"a,b\n")
    with pytest.raises(tabular.UploadError):
        tabular.detect_kind(p, name)


def test_detect_kind_rejects_extension_that_lies(tmp_path):
    exe = tmp_path / "f"
    exe.write_bytes(b"MZ\x90\x00 not a spreadsheet")
    for name in ("a.xlsx", "a.xls"):
        with pytest.raises(tabular.UploadError):
            tabular.detect_kind(exe, name)
    x = write_xlsx(tmp_path / "real.xlsx", [["a"]])
    with pytest.raises(tabular.UploadError):
        tabular.detect_kind(x, "real.csv")  # a zip renamed .csv


def test_detect_kind_rejects_macro_workbook_with_xlsx_name(tmp_path):
    p = tmp_path / "m.xlsx"
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("xl/workbook.xml", "<workbook/>")
        z.writestr("xl/vbaProject.bin", b"macro")
    with pytest.raises(tabular.UploadError, match="Macro"):
        tabular.detect_kind(p, "m.xlsx")


def test_detect_kind_rejects_zip_bomb(tmp_path):
    p = tmp_path / "bomb.xlsx"
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("xl/workbook.xml", "<workbook/>")
        z.writestr("xl/big.xml", b"0" * (50 * 1024 * 1024))
    with pytest.raises(tabular.UploadError, match="expands"):
        tabular.detect_kind(p, "bomb.xlsx")


def test_read_rows_xlsx_pads_and_drops_trailing_blank_rows(tmp_path):
    p = write_xlsx(tmp_path / "a.xlsx", [["Name", "City"], ["Apple"], [None, None]])
    rows = tabular.read_rows(p, "xlsx", None, 100)
    assert rows == [["Name", "City"], ["Apple", None]]


def test_read_rows_csv_handles_bom_and_cp1252(tmp_path):
    p = tmp_path / "a.csv"
    p.write_bytes("﻿Name\nNestlé\n".encode("utf-8"))
    assert tabular.read_rows(p, "csv", None, 10) == [["Name"], ["Nestlé"]]
    p.write_bytes("Name\nNestlé\n".encode("cp1252"))
    assert tabular.read_rows(p, "csv", None, 10) == [["Name"], ["Nestlé"]]


def test_read_rows_never_evaluates_formulas(tmp_path):
    p = write_xlsx(tmp_path / "f.xlsx", [["=1+1"]])
    # No cached value was ever computed, so the cell reads as empty, not as 2.
    assert tabular.read_rows(p, "xlsx", None, 10) == []


def test_sheet_names(tmp_path):
    p = write_xlsx(tmp_path / "a.xlsx", [["a"]], sheet="Companies")
    assert tabular.sheet_names(p, "xlsx") == ["Companies"]
    assert tabular.sheet_names(p, "csv") == []


@pytest.mark.parametrize(
    ("text", "safe"),
    [
        ("=SUM(A1)", False),
        ("+cmd|calc", False),
        ("@SUM(1)", False),
        ("-2+3", False),
        ("\tx", False),
        ("-5", True),
        ("+1,234.5", True),
        ("Apple", True),
        ("", True),
    ],
)
def test_neutralize(text, safe):
    out = tabular.neutralize(text)
    assert (out == text) is safe
    if not safe:
        assert out == "'" + text


def test_write_result_csv_roundtrip_and_injection(tmp_path):
    out = tmp_path / "out.csv"
    rows = [['=HYPERLINK("http://x")', "ok"], ["Apple", None]]
    extra = [
        ["unresolved"] + [None] * 7,
        ["resolved", "Apple", "Q312", None, None, "US", 0.9, "local"],
    ]
    tabular.write_result(out, "csv", "S", ["Name", "Note"], rows, extra)
    back = tabular.read_rows(out, "csv", None, 10)
    assert back[0][:2] == ["Name", "Note"]
    assert back[0][2:] == list(tabular.RESULT_COLUMNS)
    assert back[1][0] == '\'=HYPERLINK("http://x")'
    assert back[2][:2] == ["Apple", ""]
    assert back[2][2:5] == ["resolved", "Apple", "Q312"]


def test_write_result_xlsx_keeps_text_as_text(tmp_path):
    out = tmp_path / "out.xlsx"
    rows = [["=1+1", 5, 2.5]]
    extra = [["unresolved"] + [None] * 7]
    tabular.write_result(out, "xlsx", "S", ["a", "b", "c"], rows, extra)
    wb = openpyxl.load_workbook(out)
    cell = wb.active["A2"]
    assert cell.data_type == "s" and cell.value == "=1+1"
    assert wb.active["B2"].value == 5 and wb.active["C2"].value == 2.5


def test_xls_result_comes_back_as_xlsx_workbook(tmp_path):
    out = tmp_path / "out.xlsx"
    tabular.write_result(
        out, "xls", "Sheet: 1/2?", ["a"], [["x"]], [["blank"] + [None] * 7]
    )
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Sheet  1 2 "]  # characters Excel forbids are replaced


def test_write_result_length_mismatch_raises(tmp_path):
    with pytest.raises(ValueError):
        tabular.write_result(
            tmp_path / "o.csv",
            "csv",
            "S",
            ["a"],
            [["x"], ["y"]],
            [["blank"] + [None] * 7],
        )
